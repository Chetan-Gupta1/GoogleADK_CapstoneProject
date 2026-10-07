from pptx import Presentation
from copy import deepcopy
from pathlib import Path

from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
from renderer.ppt_utils import (
    find_shape,
    remove_shape
)


class PptRenderer:

    def __init__(self, template_path):

        self.prs = Presentation(
            template_path
        )
        self.template_slide_count = len(self.prs.slides)

    def save(self, output_path):
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(output_path))

    def get_slide(self, slide_no):
        return self.prs.slides[
            slide_no - 1
        ]

    def replace_text(
        self,
        slide_no,
        shape_name,
        text,
        bold_lead=False
    ):
        slide = self.get_slide(
            slide_no
        )
        shape = find_shape(
            slide,
            shape_name
        )
        if not shape.has_text_frame:
            raise ValueError(f"{shape_name} is not a text shape")
        frame = shape.text_frame
        paragraph_style = deepcopy(frame.paragraphs[0]._p.pPr)
        original_runs = frame.paragraphs[0].runs
        style = deepcopy(original_runs[0]._r.rPr) if original_runs else None
        font_size = original_runs[0].font.size if original_runs else None
        original_bold = original_runs[0].font.bold is True if original_runs else False
        font_family = original_runs[0].font.name or "Calibri" if original_runs else "Calibri"
        max_size = int(font_size.pt) if font_size else 18
        lines = text if isinstance(text, list) else str(text).splitlines()
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            if paragraph_style is not None:
                paragraph._p.replace(paragraph._p.get_or_add_pPr(), deepcopy(paragraph_style))
            paragraph.space_after = Pt(10 if isinstance(text, list) else 0)
            if isinstance(text, list):
                properties = paragraph._p.get_or_add_pPr()
                for child in list(properties):
                    if child.tag.endswith(("}buNone", "}buChar", "}buAutoNum")):
                        properties.remove(child)
                bullet = OxmlElement("a:buChar")
                bullet.set("char", "\u2022")
                properties.append(bullet)
                properties.set("marL", str(Pt(14)))
                properties.set("indent", str(-Pt(14)))
            run = paragraph.add_run()
            run.text = line
            if style is not None:
                run._r.replace(run._r.get_or_add_rPr(), deepcopy(style))
            run.font.italic = False
            run.font.size = Pt(max_size)
        font_filename = font_family.lower() + ("b.ttf" if original_bold else ".ttf")
        if font_family == "Cambria" and not original_bold:
            font_filename = "cambria.ttc"
        font_file = Path("C:/Windows/Fonts") / font_filename
        if font_file.exists():
            frame.fit_text(font_family=font_family, max_size=max_size, bold=original_bold, font_file=str(font_file))
        if bold_lead:
            for paragraph in frame.paragraphs:
                run = paragraph.runs[0]
                verb, separator, remainder = run.text.partition(" ")
                run.text = verb
                run.font.bold = True
                if separator:
                    tail = paragraph.add_run()
                    tail.text = separator + remainder
                    tail._r.replace(tail._r.get_or_add_rPr(), deepcopy(run._r.rPr))
                    tail.font.bold = False

    def replace_icon(
        self,
        slide_no,
        shape_name,
        icon_file
    ):

        slide = self.get_slide(
            slide_no
        )

        icon_placeholder = find_shape(
            slide,
            shape_name
        )

        left = icon_placeholder.left
        top = icon_placeholder.top
        width = icon_placeholder.width
        height = icon_placeholder.height

        with Image.open(icon_file) as image:
            image_width, image_height = image.size
        scale = min(width / image_width, height / image_height)
        picture_width, picture_height = int(image_width * scale), int(image_height * scale)
        remove_shape(icon_placeholder)
        picture = slide.shapes.add_picture(
            str(icon_file), left + (width - picture_width) // 2,
            top + (height - picture_height) // 2,
            width=picture_width, height=picture_height,
        )
        picture.name = shape_name
        self.remove_label(slide_no, shape_name + "_label")

    def remove_label(self, slide_no, shape_name):
        for shape in self.get_slide(slide_no).shapes:
            if shape.name == shape_name:
                remove_shape(shape)
                return

    def replace_chart(self, slide_no, shape_name, specification):
        slide = self.get_slide(slide_no)
        placeholder = find_shape(slide, shape_name)
        bounds = (placeholder.left, placeholder.top, placeholder.width, placeholder.height)
        chart_data = CategoryChartData()
        chart_data.categories = specification["categories"]
        for region, values in specification["series"].items():
            if len(values) != len(specification["categories"]):
                raise ValueError(f"{shape_name}: mismatched categories and values")
            chart_data.add_series(region, values)
        chart_types = {"line": XL_CHART_TYPE.LINE_MARKERS, "column": XL_CHART_TYPE.COLUMN_CLUSTERED}
        graphic_frame = slide.shapes.add_chart(
            chart_types[specification["type"]], *bounds, chart_data
        )
        graphic_frame.name = shape_name
        remove_shape(placeholder)
        self.remove_label(slide_no, shape_name + "_label")
        chart = graphic_frame.chart
        chart.has_title = True
        chart.chart_title.text_frame.text = specification["title"]
        chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)
        chart.category_axis.tick_labels.font.size = Pt(11)
        chart.category_axis.has_major_gridlines = False
        chart.category_axis.has_minor_gridlines = False
        chart.value_axis.tick_labels.font.size = Pt(11)
        chart.value_axis.minimum_scale = 0
        chart.value_axis.has_major_gridlines = False
        chart.value_axis.tick_labels.number_format = '0.0"%"' if shape_name.startswith("chart_risk") else "0"
        colors = {
            "APAC": "008878", "Europe": "BF8500",
            "Latin America": "C33B44", "North America": "3169A6",
        }
        for series in chart.series:
            series.format.line.color.rgb = RGBColor.from_string(colors.get(series.name, "555555"))
            series.format.line.width = Pt(2.5)

    def validate(self):
        if not self.template_slide_count:
            raise ValueError("The template must contain at least one slide")
        if len(self.prs.slides) != self.template_slide_count:
            raise ValueError("The final presentation must preserve the template's slide count")
        for slide in self.prs.slides:
            for shape in slide.shapes:
                if shape.name.startswith(("chart_", "icon_")) and shape.name.endswith("_label"):
                    raise ValueError(f"Instruction label remains: {shape.name}")
                if shape.has_text_frame and any(marker in shape.text for marker in ("[FILL", "[VALUE]", "Capstone Template")):
                    raise ValueError(f"Unfilled instruction remains: {shape.name}")