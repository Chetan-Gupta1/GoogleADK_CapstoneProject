from pathlib import Path
import re

from pptx import Presentation


DEFAULT_TEMPLATE = Path(__file__).resolve().parent.parent / "template" / "novaretail_template.pptx"


def inspect_template(template_path=None):
    presentation = Presentation(template_path or DEFAULT_TEMPLATE)
    if not presentation.slides:
        raise ValueError("The template must contain at least one slide")
    specifications = []
    for slide_no, slide in enumerate(presentation.slides, 1):
        specification = {
            "slide_no": slide_no, "fields": {}, "charts": {}, "icons": [],
            "subtitles": [], "kpis": {}, "footers": [], "page_numbers": [],
            "theme": "overview",
        }
        names = [shape.name for shape in slide.shapes]
        if len(names) != len(set(names)):
            raise ValueError(f"Slide {slide_no}: shape names must be unique within each slide")
        themes = set()
        for shape in slide.shapes:
            name = shape.name
            if name.startswith(("chart_", "icon_")) and name.endswith("_label"):
                continue
            field = None
            if name == "deck_title" or re.fullmatch(r"(?:deck_title|title)_\d+", name):
                field = "title"
            elif re.fullmatch(r"summary_text(?:_\d+)?", name):
                field = "summary"
            elif re.fullmatch(r"recommendations_text(?:_\d+)?", name):
                field = "recommendations"
            elif re.fullmatch(r"closing_statement(?:_\d+)?", name):
                field = "closing"
            elif name.startswith("insights_"):
                match = re.fullmatch(r"insights_(performance|customer|risk)(?:_\d+)?", name)
                if not match:
                    raise ValueError(f"Slide {slide_no}: unsupported insight theme in {name}")
                field = "insights"
                themes.add(match.group(1))
            elif name.startswith("chart_"):
                match = re.fullmatch(r"chart_(performance|customer|risk)(?:_\d+)?", name)
                if not match:
                    raise ValueError(f"Slide {slide_no}: add a dataset mapping for chart placeholder {name}")
                specification["charts"][name] = match.group(1)
                themes.add(match.group(1))
            elif re.fullmatch(r"icon_\d+", name):
                specification["icons"].append(name)
            elif re.fullmatch(r"subtitle_\d+", name):
                specification["subtitles"].append(name)
            elif re.fullmatch(r"footer_\d+", name):
                specification["footers"].append(name)
            elif re.fullmatch(r"page_num_\d+", name):
                specification["page_numbers"].append(name)
            elif re.fullmatch(r"kpi\d+_(?:value|label)(?:_\d+)?", name):
                match = re.fullmatch(r"kpi(\d+)_(value|label)(?:_\d+)?", name)
                number = int(match.group(1))
                if number not in (1, 2, 3):
                    raise ValueError(f"Slide {slide_no}: add a computed metric for {name}")
                specification["kpis"][name] = {"number": number, "kind": match.group(2)}
            elif shape.has_text_frame and any(marker in shape.text for marker in ("[FILL", "[VALUE]")):
                raise ValueError(f"Slide {slide_no}: unsupported fillable shape {name}")
            if field is not None:
                specification["fields"].setdefault(field, []).append(name)
        if len(themes) > 1:
            raise ValueError(f"Slide {slide_no}: use one data theme per slide, found {sorted(themes)}")
        if themes:
            specification["theme"] = next(iter(themes))
        elif "recommendations" in specification["fields"] or "closing" in specification["fields"]:
            specification["theme"] = "actions"
        specifications.append(specification)
    return specifications


def find_shape(slide, shape_name):

    for shape in slide.shapes:
        if shape.name == shape_name:
            return shape

    raise Exception(
        f"Shape not found: {shape_name}"
    )


def remove_shape(shape):

    sp = shape._element
    sp.getparent().remove(sp)