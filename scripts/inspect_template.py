from pathlib import Path
from pptx import Presentation
import json

BASE_DIR = Path(__file__).resolve().parent.parent

TEMPLATE_PATH = (
    BASE_DIR
    / "template"
    / "novaretail_template.pptx"
)

prs = Presentation(TEMPLATE_PATH)

all_shapes = []
for slide_no, slide in enumerate(prs.slides, start=1):

    print("\n" + "=" * 60)
    print(f"SLIDE {slide_no}")
    print("=" * 60)

    
    for shape in slide.shapes:
        print(f"\nName: {shape.name}")

        if shape.has_text_frame:
            print("TEXT:")
            print(shape.text)

        shape_info = {
            "slide": slide_no,
            "name": shape.name,
            "left": shape.left,
            "top": shape.top,
            "width": shape.width,
            "height": shape.height,
        }
        all_shapes.append(shape_info)


output_file = (
    BASE_DIR
    / "output"
    / "shape_inventory.json"
)

with open(output_file, "w") as f:
    json.dump(
        all_shapes,
        f,
        indent=4
    )

print(f"\nSaved: {output_file}")