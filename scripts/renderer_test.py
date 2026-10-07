from copy import deepcopy
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pydantic import ValidationError
from google.genai.errors import ClientError

from content.slide_generator import DeckContent, SlideGenerator, content_schema
from novaretail_agent.tools.build_presentation_tool import build_presentation
from novaretail_agent.tools.render_tool import render_presentation
from renderer.ppt_utils import find_shape, inspect_template
from renderer.ppt_renderer import PptRenderer
from scripts.analysis import run_analysis


BASE_DIR = Path(__file__).resolve().parent.parent


class PresentationTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.analysis = run_analysis()
        cls.content = SlideGenerator(cls.analysis).generate(use_ai=False)

    def test_analysis_metrics_and_row_order(self):
        kpis = self.analysis["kpis"]
        self.assertEqual(kpis["total_revenue_usd_m"], 1280)
        self.assertAlmostEqual(kpis["revenue_qoq_growth_pct"], 4.48979591836735)
        self.assertEqual(kpis["unweighted_regional_nps"], 49.75)
        self.assertEqual(self.analysis["growth_region"], "APAC")
        self.assertEqual(self.analysis["risk_region"], "Latin America")
        json.dumps(self.analysis, allow_nan=False)
        with TemporaryDirectory() as folder:
            workbook = Path(folder) / "shuffled.xlsx"
            with pd.ExcelWriter(workbook) as writer:
                for name, table in pd.read_excel(BASE_DIR / "data/novaretail_dataset.xlsx", sheet_name=None).items():
                    table.sample(frac=1, random_state=42).to_excel(writer, sheet_name=name, index=False)
            self.assertEqual(run_analysis(workbook), self.analysis)

    def test_duplicate_rows_are_rejected(self):
        with TemporaryDirectory() as folder:
            workbook = Path(folder) / "duplicate.xlsx"
            with pd.ExcelWriter(workbook) as writer:
                for name, table in pd.read_excel(BASE_DIR / "data/novaretail_dataset.xlsx", sheet_name=None).items():
                    pd.concat([table, table.iloc[[0]]]).to_excel(writer, sheet_name=name, index=False)
            with self.assertRaisesRegex(ValueError, "duplicate"):
                run_analysis(workbook)

    def test_content_schema_and_ai_evidence(self):
        with patch("llm.gemini_client.generate_text", return_value=json.dumps(self.content)) as model:
            self.assertEqual(SlideGenerator(self.analysis).generate(), self.content)
            self.assertIn("1280.0", model.call_args.args[0])
        invalid = deepcopy(self.content)
        invalid["slide_3"]["icon"] = "../untrusted.png"
        with self.assertRaises(ValidationError):
            DeckContent.model_validate(invalid)
        invalid = deepcopy(self.content)
        invalid["slide_2"]["insights"] = []
        with self.assertRaises(ValidationError):
            DeckContent.model_validate(invalid)

    def test_complete_deck_data_geometry_and_cleanup(self):
        template_path = BASE_DIR / "template/novaretail_template.pptx"
        original_bytes = template_path.read_bytes()
        template = Presentation(template_path)
        with TemporaryDirectory() as folder:
            output = render_presentation(self.content, self.analysis, Path(folder) / "review.pptx", offline=True)
            deck = Presentation(output)
        self.assertEqual(template_path.read_bytes(), original_bytes)
        specifications = inspect_template(template_path)
        self.assertEqual(len(deck.slides), len(template.slides))
        self.assertEqual(
            sum(shape.has_chart for slide in deck.slides for shape in slide.shapes),
            sum(len(specification["charts"]) for specification in specifications),
        )
        for specification in specifications:
            slide_no = specification["slide_no"]
            for name, theme in specification["charts"].items():
                chart_shape = find_shape(deck.slides[slide_no - 1], name)
                original = find_shape(template.slides[slide_no - 1], name)
                self.assertEqual(
                    (chart_shape.left, chart_shape.top, chart_shape.width, chart_shape.height),
                    (original.left, original.top, original.width, original.height),
                )
                chart_data = self.analysis["charts"][theme]
                self.assertEqual([category.label for category in chart_shape.chart.plots[0].categories], chart_data["categories"])
                self.assertEqual({series.name: list(series.values) for series in chart_shape.chart.series}, chart_data["series"])
        for slide_no, slide in enumerate(deck.slides, 1):
            picture = find_shape(slide, f"icon_{slide_no}")
            original = find_shape(template.slides[slide_no - 1], f"icon_{slide_no}")
            self.assertEqual(picture.shape_type, MSO_SHAPE_TYPE.PICTURE)
            self.assertGreaterEqual(picture.left, original.left)
            self.assertGreaterEqual(picture.top, original.top)
            self.assertLessEqual(picture.left + picture.width, original.left + original.width)
            self.assertLessEqual(picture.top + picture.height, original.top + original.height)
            with Image.open(BASE_DIR / "icons" / self.content[f"slide_{slide_no}"]["icon"]) as image:
                self.assertAlmostEqual(picture.width / picture.height, image.width / image.height, places=5)
            for shape in slide.shapes:
                self.assertFalse(shape.name.startswith(("chart_", "icon_")) and shape.name.endswith("_label"))
                if shape.has_text_frame:
                    original_shape = find_shape(template.slides[slide_no - 1], shape.name)
                    if original_shape.text_frame.paragraphs[0].runs:
                        original_color = original_shape.text_frame.paragraphs[0].runs[0].font.color
                        if original_color.type is not None:
                            self.assertEqual(shape.text_frame.paragraphs[0].runs[0].font.color.rgb, original_color.rgb)
                    for marker in ("[FILL", "[VALUE]", "Capstone Template"):
                        self.assertNotIn(marker, shape.text)
                    self.assertFalse(any(run.font.italic for paragraph in shape.text_frame.paragraphs for run in paragraph.runs))
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            self.assertTrue(run._r[0].tag.endswith("}rPr"), shape.name)
            self.assertIn("OFFLINE PREVIEW", find_shape(slide, f"footer_{slide_no}").text)
        for name, value in [("kpi1_value", "$1,280M"), ("kpi2_value", "+4.49%"), ("kpi3_value", "49.75")]:
            self.assertEqual(find_shape(deck.slides[0], name).text, value)
        for paragraph in find_shape(deck.slides[4], "recommendations_text").text_frame.paragraphs:
            self.assertTrue(paragraph.runs[0].font.bold)

    def test_build_failure_does_not_claim_success(self):
        with patch("novaretail_agent.tools.build_presentation_tool.generate_content", side_effect=RuntimeError("API unavailable")):
            result = build_presentation()
        self.assertEqual(result["status"], "error")
        self.assertEqual(result["stage"], "content")
        self.assertNotIn("file", result)

    def test_gemini_client_uses_json_schema(self):
        from llm.gemini_client import generate_text

        schema = content_schema(inspect_template())
        with patch("llm.gemini_client.genai.Client") as client, \
                patch("llm.gemini_client.load_dotenv"), \
                patch.dict(os.environ, {"GOOGLE_API_KEY": "unit-test-key", "GEMINI_MODEL": "unit-test-model"}):
            models = client.return_value.__enter__.return_value.models
            models.generate_content.return_value.text = json.dumps(self.content)
            self.assertEqual(generate_text("Test prompt", schema), json.dumps(self.content))
            config = models.generate_content.call_args.kwargs["config"]
            self.assertIsNone(config.response_schema)
            self.assertEqual(config.response_json_schema, schema.model_json_schema())
            self.assertEqual(config.response_mime_type, "application/json")
            self.assertFalse(config.response_json_schema["additionalProperties"])

    def test_api_failure_reports_safe_actionable_details(self):
        api_error = ClientError(400, {"error": {
            "code": 400, "status": "INVALID_ARGUMENT",
            "message": "Unknown additional_properties; key=unit-test-secret",
        }})
        with patch.dict(os.environ, {"GOOGLE_API_KEY": "unit-test-secret"}), \
                patch("novaretail_agent.tools.build_presentation_tool.generate_content", side_effect=api_error):
            result = build_presentation()
        self.assertEqual(result["status"], "error")
        self.assertEqual(result["stage"], "content")
        self.assertEqual(result["http_status"], 400)
        self.assertIn("additional_properties", result["message"])
        self.assertNotIn("unit-test-secret", result["message"])
        self.assertIn("[REDACTED]", result["message"])
        self.assertIn("JSON Schema", result["next_step"])
        self.assertNotIn("file", result)

    def test_renderer_preserves_variable_template_slide_count(self):
        with TemporaryDirectory() as folder:
            for count in (3, 7):
                with self.subTest(count=count):
                    template = Presentation()
                    for index in range(count):
                        template.slides.add_slide(template.slide_layouts[6])
                    path = Path(folder) / f"template_{count}.pptx"
                    template.save(path)
                    renderer = PptRenderer(path)
                    renderer.validate()
                    renderer.prs.slides.add_slide(renderer.prs.slide_layouts[6])
                    with self.assertRaisesRegex(ValueError, "preserve"):
                        renderer.validate()

    def test_changed_template_counts_roles_and_static_slides(self):
        original = Presentation(BASE_DIR / "template/novaretail_template.pptx")
        with TemporaryDirectory() as folder:
            for sources in ([4, 1, 0], [0, 1, 2, 3, 4, 1, None]):
                with self.subTest(count=len(sources)):
                    template = Presentation()
                    template.slide_width = original.slide_width
                    template.slide_height = original.slide_height
                    for source_no in sources:
                        slide = template.slides.add_slide(template.slide_layouts[6])
                        if source_no is None:
                            note = slide.shapes.add_textbox(100000, 100000, 4000000, 500000)
                            note.name = "static_note"
                            note.text = "Appendix: data definitions"
                        else:
                            for shape in original.slides[source_no].shapes:
                                slide.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
                    if len(sources) == 7:
                        for shape in template.slides[5].shapes:
                            replacements = {
                                "chart_performance": "chart_performance_6",
                                "chart_performance_label": "chart_performance_6_label",
                                "insights_performance": "insights_performance_6",
                                "title_2": "title_6", "icon_2": "icon_6",
                                "icon_2_label": "icon_6_label", "footer_2": "footer_6",
                                "page_num_2": "page_num_6",
                            }
                            shape.name = replacements.get(shape.name, shape.name)
                    template_path = Path(folder) / f"template_{len(sources)}.pptx"
                    template.save(template_path)
                    generator = SlideGenerator(self.analysis, template_path=template_path)
                    content = generator.generate(use_ai=False)
                    self.assertEqual(set(content), {f"slide_{number}" for number in range(1, len(sources) + 1)})
                    with patch("llm.gemini_client.generate_text", return_value=json.dumps(content)) as model:
                        self.assertEqual(generator.generate(), content)
                        self.assertIn("Template specification", model.call_args.args[0])
                        schema = model.call_args.args[1].model_json_schema()
                        self.assertEqual(set(schema["required"]), set(content))
                    output = render_presentation(
                        content, self.analysis, Path(folder) / f"review_{len(sources)}.pptx",
                        offline=True, template_path=template_path,
                    )
                    deck = Presentation(output)
                    self.assertEqual(len(deck.slides), len(sources))
                    for specification in inspect_template(template_path):
                        slide_no = specification["slide_no"]
                        slide = deck.slides[slide_no - 1]
                        for field, names in specification["fields"].items():
                            for name in names:
                                self.assertEqual(find_shape(slide, name).text.replace("\n", " "),
                                                 " ".join(content[f"slide_{slide_no}"][field])
                                                 if isinstance(content[f"slide_{slide_no}"][field], list)
                                                 else content[f"slide_{slide_no}"][field])
                        for name, theme in specification["charts"].items():
                            chart = find_shape(slide, name)
                            placeholder = find_shape(template.slides[slide_no - 1], name)
                            self.assertEqual((chart.left, chart.top, chart.width, chart.height),
                                             (placeholder.left, placeholder.top, placeholder.width, placeholder.height))
                            self.assertEqual({series.name: list(series.values) for series in chart.chart.series},
                                             self.analysis["charts"][theme]["series"])
                        for name in specification["icons"]:
                            self.assertEqual(find_shape(slide, name).shape_type, MSO_SHAPE_TYPE.PICTURE)
                        for name in specification["page_numbers"]:
                            self.assertEqual(find_shape(slide, name).text, str(slide_no))
                        self.assertFalse(any(shape.name.endswith("_label") and shape.name.startswith(("chart_", "icon_")) for shape in slide.shapes))
                    if len(sources) == 3:
                        self.assertIn("recommendations", content["slide_1"])
                        self.assertEqual(find_shape(deck.slides[2], "kpi1_value").text, "$1,280M")
                    else:
                        self.assertEqual(content["slide_7"], {})
                        self.assertEqual(find_shape(deck.slides[6], "static_note").text, "Appendix: data definitions")

    def test_content_must_match_template_contract(self):
        schema = content_schema(inspect_template())
        missing = deepcopy(self.content)
        missing.pop(next(iter(missing)))
        with self.assertRaises(ValidationError):
            schema.model_validate(missing)
        extra = deepcopy(self.content)
        extra[f"slide_{len(extra) + 1}"] = {}
        with self.assertRaises(ValidationError):
            schema.model_validate(extra)


if __name__ == "__main__":
    unittest.main()